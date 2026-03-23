"""原稿パーサー: docx/xlsx/txt形式の原稿を読み込み、番号付きエントリのリストを返す。"""

import re
from pathlib import Path


def parse_manuscript(filepath: str | Path) -> list[dict]:
    """原稿ファイルを読み込み、エントリのリストを返す。

    Args:
        filepath: 原稿ファイルパス (.docx / .xlsx / .txt)

    Returns:
        [{"id": "1", "text": "..."}, ...]
    """
    filepath = Path(filepath)
    suffix = filepath.suffix.lower()

    if suffix == ".docx":
        text = _read_docx(filepath)
    elif suffix == ".xlsx":
        return _read_xlsx(filepath)
    elif suffix == ".txt":
        text = filepath.read_text(encoding="utf-8")
    else:
        raise ValueError(f"未対応のファイル形式: {suffix}")

    return _parse_numbered_text(text)


def _read_docx(filepath: Path) -> str:
    """docxファイルからテキストを抽出する。"""
    import docx

    doc = docx.Document(str(filepath))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_xlsx(filepath: Path) -> list[dict]:
    """xlsxファイルからエントリを読み取る。"""
    import openpyxl

    wb = openpyxl.load_workbook(str(filepath), read_only=True)
    ws = wb.active

    rows = []
    for row in ws.iter_rows(values_only=True):
        non_none = [c for c in row if c is not None]
        if non_none:
            rows.append(non_none)
    wb.close()

    if not rows:
        return []

    # 2列以上ある場合: 1列目=番号、2列目=テキスト
    if len(rows[0]) >= 2:
        return [{"id": str(row[0]), "text": str(row[1])} for row in rows]

    # 1列のみの場合: テキストを結合して番号検出ロジックを適用
    text = "\n".join(str(row[0]) for row in rows)
    return _parse_numbered_text(text)


def _parse_numbered_text(text: str) -> list[dict]:
    """テキストを番号付きエントリに分割する。

    番号行（数字のみの行）を検出し、次の番号行までのテキストを1エントリとする。
    番号行がなければ空行で区切られたブロックを1エントリとする。
    """
    lines = text.split("\n")

    # 番号行の検出
    number_indices = []
    for i, line in enumerate(lines):
        if re.match(r"^\d+$", line.strip()):
            number_indices.append(i)

    if number_indices:
        return _parse_with_numbers(lines, number_indices)
    else:
        return _parse_without_numbers(lines)


def _parse_with_numbers(lines: list[str], number_indices: list[int]) -> list[dict]:
    """番号行で区切られたエントリを生成する。"""
    entries = []
    for idx, num_line_idx in enumerate(number_indices):
        entry_id = lines[num_line_idx].strip()
        # 次の番号行までの範囲を取得
        if idx + 1 < len(number_indices):
            end = number_indices[idx + 1]
        else:
            end = len(lines)
        # 番号行の次の行からテキストを収集
        text_lines = []
        for line in lines[num_line_idx + 1 : end]:
            stripped = line.strip()
            if stripped:
                text_lines.append(stripped)
        if text_lines:
            entries.append({"id": entry_id, "text": "\n".join(text_lines)})
    return entries


def _parse_without_numbers(lines: list[str]) -> list[dict]:
    """空行で区切られたブロックを1エントリとする。"""
    entries = []
    current_block = []
    for line in lines:
        if line.strip():
            current_block.append(line.strip())
        else:
            if current_block:
                entries.append({"id": str(len(entries) + 1), "text": "\n".join(current_block), "auto_numbered": True})
                current_block = []
    if current_block:
        entries.append({"id": str(len(entries) + 1), "text": "\n".join(current_block), "auto_numbered": True})
    return entries
